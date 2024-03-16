
from PyQt5 import  QtCore, QtGui, QtWidgets
from PyQt5.QtWidgets import QDialog, QTextEdit, QScrollArea, QVBoxLayout, QPushButton, QLabel, QApplication
import os
from openpyxl import Workbook
from docx import Document
from create_data_base import Data_base_meth
from datetime import date


class Scroll_move(QDialog):
    def __init__(self):
        super().__init__()
        self.my_db = Data_base_meth(database_name="new_bd.db")
        self.setupUi(self)
    def setupUi(self, Dialog):
        Dialog.setObjectName("Dialog")
        Dialog.resize(446, 346)
        self.pushButton_5 = QtWidgets.QPushButton(Dialog)
        self.pushButton_5.setGeometry(QtCore.QRect(0, 300, 241, 32))
        self.pushButton_5.setObjectName("pushButton_5")
        self.label = QtWidgets.QLabel(Dialog)
        self.label.setGeometry(QtCore.QRect(10, 0, 125, 16))
        self.label.setObjectName("label")
        self.layoutWidget_2 = QtWidgets.QWidget(Dialog)
        self.layoutWidget_2.setGeometry(QtCore.QRect(10, 30, 91, 251))
        self.layoutWidget_2.setObjectName("layoutWidget_2")
        self.verticalLayout_2 = QtWidgets.QVBoxLayout(self.layoutWidget_2)
        self.verticalLayout_2.setContentsMargins(0, 0, 0, 0)
        self.verticalLayout_2.setObjectName("verticalLayout_2")
        self.label_3 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_3.setObjectName("label_3")
        self.verticalLayout_2.addWidget(self.label_3)
        self.label_2 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_2.setObjectName("label_2")
        self.verticalLayout_2.addWidget(self.label_2)
        self.label_5 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_5.setObjectName("label_5")
        self.verticalLayout_2.addWidget(self.label_5)
        self.label_7 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_7.setObjectName("label_7")
        self.verticalLayout_2.addWidget(self.label_7)
        self.label_6 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_6.setObjectName("label_6")
        self.verticalLayout_2.addWidget(self.label_6)
        self.label_4 = QtWidgets.QLabel(self.layoutWidget_2)
        self.label_4.setObjectName("label_4")
        self.verticalLayout_2.addWidget(self.label_4)
        self.scrollArea = QtWidgets.QScrollArea(Dialog)
        self.scrollArea.setGeometry(QtCore.QRect(110, 30, 301, 31))
        self.scrollArea.setWidgetResizable(True)
        self.scrollArea.setObjectName("scrollArea")
        self.scrollAreaWidgetContents = QtWidgets.QWidget()
        self.scrollAreaWidgetContents.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents.setObjectName("scrollAreaWidgetContents")
        self.listWidget = QtWidgets.QListWidget(Dialog)
        self.listWidget.setGeometry(QtCore.QRect(110, 30, 299, 29))
        self.listWidget.setObjectName("listWidget")
        data = self.my_db.select_all_from_db()
        for item in data:
            name = str(item[1])
            list_item = QtWidgets.QListWidgetItem(name)
            self.listWidget.addItem(list_item)
        self.listWidget.itemClicked.connect(self.save_data)
        self.textEdit = QtWidgets.QTextEdit(self.scrollAreaWidgetContents)
        self.textEdit.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit.setObjectName("textEdit")
        self.scrollArea.setWidget(self.scrollAreaWidgetContents)
        self.scrollArea_2 = QtWidgets.QScrollArea(Dialog)
        self.scrollArea_2.setGeometry(QtCore.QRect(110, 70, 301, 31))
        self.scrollArea_2.setWidgetResizable(True)
        self.scrollArea_2.setObjectName("scrollArea_2")
        self.scrollAreaWidgetContents_2 = QtWidgets.QWidget()
        self.scrollAreaWidgetContents_2.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents_2.setObjectName("scrollAreaWidgetContents_2")
        self.textEdit_2 = QtWidgets.QTextEdit(self.scrollAreaWidgetContents_2)
        self.textEdit_2.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit_2.setObjectName("textEdit_2")
        self.scrollArea_2.setWidget(self.scrollAreaWidgetContents_2)
        self.scrollArea_3 = QtWidgets.QScrollArea(Dialog)
        self.scrollArea_3.setGeometry(QtCore.QRect(110, 120, 301, 31))
        self.scrollArea_3.setWidgetResizable(True)
        self.scrollArea_3.setObjectName("scrollArea_3")
        self.scrollAreaWidgetContents_3 = QtWidgets.QWidget()
        self.scrollAreaWidgetContents_3.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents_3.setObjectName("scrollAreaWidgetContents_3")
        self.textEdit_3 = QtWidgets.QTextEdit(self.scrollAreaWidgetContents_3)
        self.textEdit_3.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit_3.setObjectName("textEdit_3")
        self.scrollArea_3.setWidget(self.scrollAreaWidgetContents_3)
        self.scrollArea_4 = QtWidgets.QScrollArea(Dialog)
        self.scrollArea_4.setGeometry(QtCore.QRect(110, 160, 301, 31))
        self.scrollArea_4.setWidgetResizable(True)
        self.scrollArea_4.setObjectName("scrollArea_4")
        self.scrollAreaWidgetContents_4 = QtWidgets.QWidget()
        self.scrollAreaWidgetContents_4.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents_4.setObjectName("scrollAreaWidgetContents_4")
        self.textEdit_4 = QtWidgets.QTextEdit(self.scrollAreaWidgetContents_4)
        self.textEdit_4.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit_4.setObjectName("textEdit_4")
        self.scrollArea_4.setWidget(self.scrollAreaWidgetContents_4)
        self.scrollArea_5 = QtWidgets.QScrollArea(Dialog)
        self.scrollArea_5.setGeometry(QtCore.QRect(110, 210, 301, 31))
        self.scrollArea_5.setWidgetResizable(True)
        self.scrollArea_5.setObjectName("scrollArea_5")
        self.scrollAreaWidgetContents_5 = QtWidgets.QWidget()
        self.scrollAreaWidgetContents_5.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents_5.setObjectName("scrollAreaWidgetContents_5")
        self.textEdit_5 = QtWidgets.QTextEdit(self.scrollAreaWidgetContents_5)
        self.textEdit_5.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit_5.setObjectName("textEdit_5")
        self.scrollArea_5.setWidget(self.scrollAreaWidgetContents_5)
        self.scrollArea_6 = QtWidgets.QScrollArea(Dialog)
        self.scrollArea_6.setGeometry(QtCore.QRect(110, 250, 301, 31))
        self.scrollArea_6.setWidgetResizable(True)
        self.scrollArea_6.setObjectName("scrollArea_6")

        self.scrollAreaWidgetContents_6 = QtWidgets.QWidget()
        self.scrollAreaWidgetContents_6.setGeometry(QtCore.QRect(0, 0, 299, 29))
        self.scrollAreaWidgetContents_6.setObjectName("scrollAreaWidgetContents_6")
        self.textEdit_6 = QtWidgets.QTextEdit(self.scrollAreaWidgetContents_6)
        self.textEdit_6.setGeometry(QtCore.QRect(0, 0, 301, 31))
        self.textEdit_6.setObjectName("textEdit_6")
        today = date.today()
        self.textEdit_6.setPlainText(str(today))
        self.scrollArea_6.setWidget(self.scrollAreaWidgetContents_6)
        # self.textEdit.setPlainText('\n'.join(str(item[1]) for item in create_data_base.select_all_from_db()))
        self.pushButton_5.clicked.connect(self.save_data)
        self.retranslateUi(Dialog)
        QtCore.QMetaObject.connectSlotsByName(Dialog)

    def save_data1(self):
        print(1111)
    def save_data(self):
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
        current_item = self.listWidget.currentItem()

        name_text = current_item.text()

        data = {
            "Данные водителя": name_text,
            "Принимающий": self.textEdit_2.toPlainText(),
            "Склад №А": self.textEdit_3.toPlainText(),
            "Склад №Б": self.textEdit_4.toPlainText(),
            "Товар": self.textEdit_5.toPlainText(),
            "Дата": self.textEdit_6.toPlainText()
        }

        header = ["Данные водителя", "Принимающий", "Склад №А", "Склад №Б","Товар", "Дата"]


        excel_file = os.path.join(desktop_path, "move_data.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(header)
        row_data = [data[key] for key in header]
        ws.append(row_data)
        wb.save(excel_file)


        word_file = os.path.join(desktop_path, "move_data.docx")
        doc = Document()
        for key, value in data.items():
            doc.add_paragraph(f"{key}: {value}")
        doc.save(word_file)

    def item_selected(self, item):

        selected_text = item.text()
        print("Selected:", selected_text)
    def retranslateUi(self, Dialog):
        _translate = QtCore.QCoreApplication.translate
        Dialog.setWindowTitle(_translate("Dialog", "Dialog"))
        self.pushButton_5.setText(_translate("Dialog", "Печать документов"))
        self.label.setText(_translate("Dialog", "Передвижение"))
        self.label_3.setText(_translate("Dialog", "Данные водителя"))
        self.label_2.setText(_translate("Dialog", "Принимающий"))
        self.label_5.setText(_translate("Dialog", "Склад №А"))
        self.label_7.setText(_translate("Dialog", "Склад №Б"))
        self.label_6.setText(_translate("Dialog", "Товар"))
        self.label_4.setText(_translate("Dialog", "Дата"))


if __name__ == "__main__":
    import sys

    app = QApplication(sys.argv)
    dialog = Scroll_move()
    dialog.show()
    sys.exit(app.exec_())
